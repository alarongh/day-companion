from __future__ import annotations

import csv
import hashlib
import io
import json
import re
import shutil
import subprocess
import tempfile
import zipfile
from datetime import date, datetime, timedelta
from pathlib import Path
from zoneinfo import ZoneInfo

try:
    import pymupdf
except ImportError:
    pymupdf = None
import recurring_ical_events
from icalendar import Calendar
from openpyxl import load_workbook

from .domain import Lesson, Profile

MAX_LESSONS = 5000
FILE_ERRORS = (zipfile.BadZipFile, subprocess.SubprocessError, OSError, EOFError)
if pymupdf is not None:
    FILE_ERRORS += (pymupdf.FileDataError,)
ALIASES = {
    "дата": "date",
    "день": "weekday",
    "день недели": "weekday",
    "начало": "start",
    "конец": "end",
    "окончание": "end",
    "предмет": "title",
    "название": "title",
    "аудитория": "room",
    "корпус": "campus",
    "подгруппа": "subgroup",
    "недели": "weeks",
    "чётность": "parity",
    "четность": "parity",
    "действует с": "valid_from",
    "действует до": "valid_until",
}
DAYS = {
    name: i
    for i, names in enumerate(
        [
            ("пн", "понедельник"),
            ("вт", "вторник"),
            ("ср", "среда"),
            ("чт", "четверг"),
            ("пт", "пятница"),
            ("сб", "суббота"),
            ("вс", "воскресенье"),
        ]
    )
    for name in names
}


def decode(data: bytes) -> str:
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("cp1251")


def validate_rows(rows: list[dict]) -> list[Lesson]:
    if len(rows) > MAX_LESSONS:
        raise ValueError(f"Слишком много занятий: максимум {MAX_LESSONS}")
    lessons = []
    for index, raw in enumerate(rows, 1):
        row = {}
        for key, value in raw.items():
            key = ALIASES.get(str(key).strip().lower(), str(key).strip().lower())
            if value is not None and value != "":
                row[key] = value
        for key in ("date", "valid_from", "valid_until"):
            if key in row:
                value = row[key]
                if isinstance(value, datetime):
                    value = value.date()
                elif isinstance(value, str) and re.fullmatch(r"\d{2}\.\d{2}\.\d{4}", value.strip()):
                    value = datetime.strptime(value.strip(), "%d.%m.%Y").date()
                row[key] = value
        if "weekday" in row:
            value = str(row["weekday"]).strip().lower()
            row["weekday"] = DAYS.get(value, row["weekday"])
        for key in ("start", "end"):
            if key in row:
                row[key] = (
                    row[key].strftime("%H:%M") if hasattr(row[key], "strftime") else str(row[key]).strip()
                )
        for key in ("room", "campus", "title", "subgroup"):
            if key in row:
                row[key] = str(row[key]).strip()
        if "weeks" in row and not isinstance(row["weeks"], list):
            row["weeks"] = [int(x.strip()) for x in str(row["weeks"]).split(",") if x.strip()]
        if "parity" in row:
            row["parity"] = {"все": "all", "чет": "even", "чёт": "even", "нечет": "odd", "нечёт": "odd"}.get(
                str(row["parity"]).lower(), row["parity"]
            )
        try:
            lessons.append(Lesson.model_validate(row))
        except (ValueError, TypeError) as exc:
            raise ValueError(f"Строка {index}: {exc}") from exc
    if not lessons:
        raise ValueError("В файле не найдено занятий")
    unique = {x.id: x for x in lessons}
    if len(unique) != len(lessons):
        raise ValueError("В файле есть одинаковые ID или повторяющиеся занятия. Уберите дубли.")
    return lessons


def parse_text(text: str) -> list[Lesson]:
    text = text.strip()
    if text.startswith(("[", "{")):
        value = json.loads(text)
        if isinstance(value, dict):
            value = value.get("lessons", [])
        return validate_rows(value)
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    if not lines:
        raise ValueError("Файл пуст")
    if any(k in lines[0].lower() for k in ("weekday", "date", "предмет", "день", "дата")) and any(
        s in lines[0] for s in (";", "\t", ",")
    ):
        dialect = csv.Sniffer().sniff("\n".join(lines[:10]), delimiters=";,\t")
        return validate_rows(list(csv.DictReader(io.StringIO(text), dialect=dialect)))
    rows = []
    for index, line in enumerate(lines, 1):
        if line.startswith("#"):
            continue
        # Every nonempty line must parse. Never silently discard ambiguous schedule rows.
        pieces = [x.strip() for x in line.split("|")]
        if len(pieces) < 3:
            raise ValueError(
                f"Строка {index} не распознана. Формат: Пн | 09:00-10:30 | Предмет | Аудитория | Корпус | все | Подгруппа"
            )
        match = re.fullmatch(r"(\d{1,2}:\d{2})\s*[-–—]\s*(\d{1,2}:\d{2})", pieces[1])
        if not match:
            raise ValueError(f"Строка {index}: не удалось прочитать время")
        selector = "weekday" if pieces[0].lower() in DAYS else "date"
        row = {selector: pieces[0], "start": match[1], "end": match[2], "title": pieces[2]}
        for key, value in zip(("room", "campus", "parity", "subgroup"), pieces[3:], strict=False):
            row[key] = value
        rows.append(row)
    return validate_rows(rows)


def parse_ics(data: bytes, profile: Profile, first: date, last: date) -> list[Lesson]:
    # Some upstream calendars use CRCRLF; normalize before unfolding.
    data = data.replace(b"\r\r\n", b"\r\n")
    cal = Calendar.from_ical(data)
    zone = ZoneInfo(profile.timezone)
    begin = datetime.combine(first, datetime.min.time(), zone)
    finish = datetime.combine(last + timedelta(days=1), datetime.min.time(), zone)
    events = recurring_ical_events.of(cal).between(begin, finish)
    rows = []
    for event in events:
        if str(event.get("STATUS", "")).upper() == "CANCELLED":
            continue
        start = event.decoded("DTSTART", None)
        end = event.decoded("DTEND", None)
        if not isinstance(start, datetime):
            continue  # All-day university week markers are not lessons.
        if end is None and event.get("DURATION"):
            end = start + event.decoded("DURATION")
        if not isinstance(end, datetime):
            raise ValueError("В календаре есть занятие без времени окончания")
        start = start.replace(tzinfo=zone) if start.tzinfo is None else start.astimezone(zone)
        end = end.replace(tzinfo=zone) if end.tzinfo is None else end.astimezone(zone)
        if not first <= start.date() <= last:
            continue
        if start.date() != end.date():
            raise ValueError("Занятия через полночь пока не поддерживаются")
        auditorium = event.get("X-META-AUDITORIUM")
        room, campus = str(event.get("LOCATION", "")), ""
        if auditorium is not None:
            room = str(auditorium.params.get("NUMBER", auditorium))
            campus = str(auditorium.params.get("CAMPUS", ""))
        else:
            m = re.fullmatch(r"(.+?)\s*\(([^)]+)\)", room)
            if m:
                room, campus = m.groups()
        uid = str(event.get("UID", ""))
        stable_id = hashlib.sha256(f"{uid}:{start.date()}".encode()).hexdigest()[:16] if uid else ""
        rows.append(
            dict(
                id=stable_id,
                date=start.date(),
                start=start.strftime("%H:%M"),
                end=end.strftime("%H:%M"),
                title=str(event.get("X-META-DISCIPLINE", event.get("SUMMARY", "Занятие"))),
                room=room,
                campus=campus,
                subgroup=str(event.get("X-META-SUBGROUP", "")),
            )
        )
        if len(rows) > MAX_LESSONS:
            raise ValueError("Слишком много событий в календаре")
    return validate_rows(rows) if rows else []


def ocr_image(data: bytes, suffix: str, tesseract: str = "") -> str:
    with tempfile.TemporaryDirectory(prefix="companion-ocr-") as directory:
        source = Path(directory) / f"page{suffix}"
        source.write_bytes(data)
        command = tesseract or shutil.which("tesseract")
        if command:
            result = subprocess.run(
                [command, str(source), "stdout", "-l", "rus+eng"], capture_output=True, timeout=60
            )
            if result.returncode:
                raise ValueError("Локальный OCR не смог прочитать изображение")
            return result.stdout.decode("utf-8", "replace")
        if shutil.which("powershell.exe"):
            script = Path(__file__).parent.parent / "scripts" / "ocr-windows.ps1"
            result = subprocess.run(
                ["powershell.exe", "-NoProfile", "-File", str(script), str(source)],
                capture_output=True,
                timeout=60,
            )
            if result.returncode == 0 and result.stdout.strip():
                return result.stdout.decode("utf-8-sig", "replace")
        raise ValueError(
            "Локальное распознавание недоступно. Пришли ICS/CSV/Excel/JSON или текст. Для OCR на сервере установи Tesseract с языками rus+eng."
        )


def import_file(name: str, data: bytes, profile: Profile, today: date, tesseract="") -> list[Lesson]:
    try:
        return _import_file(name, data, profile, today, tesseract)
    except FILE_ERRORS as exc:
        raise ValueError(
            "Не удалось прочитать файл или распознать текст. Проверь файл и попробуй ICS/CSV/Excel/JSON или текстовый шаблон."
        ) from exc


def _import_file(name: str, data: bytes, profile: Profile, today: date, tesseract="") -> list[Lesson]:
    if len(data) > 10 * 1024 * 1024:
        raise ValueError("Максимальный размер файла 10 МБ")
    suffix = Path(name).suffix.lower()
    if suffix == ".ics":
        return parse_ics(data, profile, today - timedelta(days=7), today + timedelta(days=180))
    if suffix in (".csv", ".txt", ".json"):
        return parse_text(decode(data))
    if suffix == ".xlsx":
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(x.file_size for x in archive.infolist()) > 50 * 1024 * 1024:
                raise ValueError("Распакованный Excel превышает 50 МБ")
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            sheet = book.active
            if sheet.max_row > 5001 or sheet.max_column > 30:
                raise ValueError("Excel: максимум 5000 строк и 30 столбцов")
            iterator = sheet.iter_rows(values_only=True)
            header = next(iterator, ())
            rows = [
                dict(zip(header, row, strict=True)) for row in iterator if any(v is not None for v in row)
            ]
            return validate_rows(rows)
        finally:
            book.close()
    if suffix in (".png", ".jpg", ".jpeg", ".webp"):
        if pymupdf is None:
            raise ValueError(
                "В бесплатном облаке OCR недоступен. Передай фото своему ИИ вместе с промптом /config и верни JSON/TXT, либо загрузи ICS/CSV/XLSX."
            )
        return parse_text(ocr_image(data, suffix, tesseract))
    if suffix == ".pdf":
        if pymupdf is None:
            raise ValueError(
                "В бесплатном облаке обработка PDF недоступна. Передай PDF своему ИИ вместе с промптом /config и верни JSON/TXT, либо загрузи ICS/CSV/XLSX."
            )
        with pymupdf.open(stream=data, filetype="pdf") as document:
            if len(document) > 15:
                raise ValueError("PDF: максимум 15 страниц")
            parts = []
            for page in document:
                text = page.get_text(sort=True)
                if len(text.strip()) < 30:
                    pix = page.get_pixmap(dpi=180)
                    text = ocr_image(pix.tobytes("png"), ".png", tesseract)
                parts.append(text)
        return parse_text("\n".join(parts))
    raise ValueError(
        "Поддерживаются ICS, CSV, XLSX, JSON, TXT, PDF и изображения. Табличные PDF/скриншоты со сложной вёрсткой могут потребовать преобразования в шаблон."
    )
